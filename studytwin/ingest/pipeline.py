"""Ingestion pipeline: any supported file (or a YouTube URL) -> citable knowledge units.

    PDF textbook  -> paragraph units (page) + figure units (page + bbox, vision description)
    PDF slides    -> one unit per slide (slide number) + diagram units (vision description)
    PPTX          -> one unit per slide (+ notes) + picture units
    Video/audio   -> Whisper transcript -> ~60 s units (start/end timestamps)
    SRT/VTT       -> caption units (timestamps)
    YouTube URL   -> published captions -> units (timestamps)

No manual preprocessing: the file type decides the path, and every unit stores the location
(page / slide / timestamp) that the tutor cites and the viewer opens.
"""

from __future__ import annotations

import re
import shutil
import time
from pathlib import Path
from typing import Any, Callable

from .. import retrieval, store
from ..config import settings
from . import figures
from .pdf import RawUnit, ingest_pdf
from .video import (
    fetch_youtube_transcript,
    fmt_ts,
    parse_caption_file,
    transcribe_media,
    window_segments,
    youtube_id,
)

VIDEO_EXT = {".mp4", ".mkv", ".webm", ".mov", ".avi", ".m4v"}
AUDIO_EXT = {".mp3", ".wav", ".m4a", ".ogg", ".flac"}
CAPTION_EXT = {".srt", ".vtt"}
SUPPORTED = {".pdf", ".pptx"} | VIDEO_EXT | AUDIO_EXT | CAPTION_EXT

Progress = Callable[[float, str], None]


def est_tokens(text: str) -> int:
    return max(1, round(len(text) / 4))


def _noop(_: float, __: str) -> None:
    return None


def ingest_pptx(path: Path, out_dir: Path, prefix: str, short_title: str) -> list[RawUnit]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(path))
    units: list[RawUnit] = []
    out_dir.mkdir(parents=True, exist_ok=True)
    for i, slide in enumerate(prs.slides, start=1):
        texts: list[str] = []
        pics: list[Path] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                img = shape.image
                p = out_dir / f"slide_{i:03d}_pic{len(pics) + 1}.{img.ext}"
                p.write_bytes(img.blob)
                pics.append(p)
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        body = "\n".join(texts) + (f"\nSpeaker notes: {notes}" if notes else "")
        title = texts[0].split("\n")[0][:90] if texts else f"Slide {i}"
        if body.strip():
            units.append(RawUnit(kind="slide", text=body, loc={"slide": i}, label=f"{short_title} · slide {i}",
                                 section=title, image=f"{prefix}/{pics[0].name}" if pics else None))
        for pi, p in enumerate(pics, start=1):
            units.append(RawUnit(kind="figure", text=f"Picture on slide {i}: {title}", loc={"slide": i, "picture": pi},
                                 label=f"{short_title} · slide {i} · figure", section=title, image=f"{prefix}/{p.name}"))
    return units


def _enrich_visuals(units: list[RawUnit], source_kind: str, progress: Progress) -> list[RawUnit]:
    """Vision pass: slide transcription + diagram descriptions; figure descriptions."""
    jobs, targets = [], []
    for u in units:
        if u.kind == "slide" and u.image:
            jobs.append((settings.media_dir / u.image, "one lecture slide", f"Slide title: {u.section}. Extracted text: {u.text[:400]}"))
            targets.append(u)
        elif u.kind == "figure" and u.image:
            jobs.append((settings.media_dir / u.image, "a figure cropped from a textbook page", u.text[:400]))
            targets.append(u)
    if not jobs:
        return units
    progress(0.55, f"Describing {len(jobs)} slides/figures with a vision model")
    results = figures.describe_many(jobs)
    extra: list[RawUnit] = []
    for u, res in zip(targets, results):
        if not res or res.get("error"):
            continue
        if u.kind == "slide":
            if res.get("text"):
                u.extra["pdf_text"] = u.text
                u.text = res["text"]
            if res.get("has_diagram") and res.get("description"):
                extra.append(RawUnit(
                    kind="figure",
                    text=f"Diagram on {u.label} ({u.section}): {res['description']}",
                    loc=dict(u.loc), label=f"{u.label} · diagram", section=u.section,
                    image=u.image, page_image=u.page_image, extra={"vision": True},
                ))
        else:
            parts = [u.text]
            if res.get("description"):
                parts.append(f"Visual description: {res['description']}")
            if res.get("text"):
                parts.append(f"Text in figure: {res['text']}")
            u.text = "\n".join(parts)
            u.extra["vision"] = True
    # keep diagram units next to their slide
    if extra:
        by_slide = {(e.loc.get("slide"), e.loc.get("page")): e for e in extra}
        merged: list[RawUnit] = []
        for u in units:
            merged.append(u)
            if u.kind == "slide":
                e = by_slide.get((u.loc.get("slide"), u.loc.get("page")))
                if e:
                    merged.append(e)
        return merged
    return units


def ingest_source(
    path: str | Path | None = None,
    *,
    youtube_url: str | None = None,
    kind: str | None = None,
    title: str | None = None,
    short_title: str | None = None,
    license: str = "",
    attribution: str = "",
    page_offset: int = 0,
    remote_url: str = "",
    copy_file: bool = True,
    enrich: bool = True,
    progress: Progress = _noop,
) -> str:
    """Ingest one source and return its id. Raises ValueError for unsupported input."""
    settings.ensure_dirs()
    sid = store.new_id("src_")
    meta: dict[str, Any] = {}
    t0 = time.perf_counter()

    if youtube_url:
        vid = youtube_id(youtube_url)
        if not vid:
            raise ValueError("Not a YouTube URL")
        progress(0.1, "Fetching YouTube captions")
        segs = fetch_youtube_transcript(vid)
        kind, title = "video", title or f"YouTube {vid}"
        short_title = short_title or "Video"
        remote_url = f"https://www.youtube.com/watch?v={vid}"
        meta.update({"youtube_id": vid, "transcript": "youtube-captions"})
        raw_units = _video_units(segs, short_title)
        src_path, filename = "", ""
    else:
        p = Path(path) if path else None
        if p is None or not p.exists():
            raise ValueError(f"File not found: {path}")
        ext = p.suffix.lower()
        if ext not in SUPPORTED:
            raise ValueError(f"Unsupported file type {ext}. Supported: {', '.join(sorted(SUPPORTED))}")
        filename = p.name
        if copy_file:
            dst_dir = settings.uploads_dir / sid
            dst_dir.mkdir(parents=True, exist_ok=True)
            dst = dst_dir / re.sub(r"[^A-Za-z0-9._-]", "_", p.name)
            if p.resolve() != dst.resolve():
                shutil.copyfile(p, dst)
            p = dst
        src_path = str(p)
        title = title or Path(filename).stem.replace("_", " ")
        if ext == ".pdf":
            progress(0.1, "Reading PDF pages")
            kind, raw_units = ingest_pdf(p, settings.media_dir, sid, kind=kind if kind in ("slides", "textbook") else None,
                                         page_offset=page_offset, short_title=short_title or title[:30])
            short_title = short_title or ("Slides" if kind == "slides" else "Textbook")
        elif ext == ".pptx":
            progress(0.1, "Reading PowerPoint slides")
            kind = "slides"
            raw_units = ingest_pptx(p, settings.media_dir / sid, sid, short_title or title[:30])
        elif ext in CAPTION_EXT:
            progress(0.2, "Parsing captions")
            kind = "video"
            segs = parse_caption_file(p)
            meta["transcript"] = "caption-file"
            raw_units = _video_units(segs, short_title or title[:30])
        else:
            kind = "video"

            def wprog(f: float, m: str) -> None:
                progress(0.05 + 0.75 * f, m)

            segs, tmeta = transcribe_media(p, progress=wprog)
            meta.update({"transcript": "whisper", **tmeta})
            raw_units = _video_units(segs, short_title or title[:30])

    if enrich and kind in ("slides", "textbook"):
        raw_units = _enrich_visuals(raw_units, kind, progress)

    progress(0.85, f"Storing {len(raw_units)} knowledge units")
    store.insert("sources", {
        "id": sid, "kind": kind, "title": title, "filename": filename, "path": src_path,
        "remote_url": remote_url, "license": license, "attribution": attribution,
        "meta": {**meta, "short_title": short_title, "page_offset": page_offset, "units": len(raw_units)},
        "created_at": time.time(),
    })
    raw_units = [u for u in raw_units if u.kind != "text" or len(u.text.split()) >= 8]  # drop orphan headings
    rows = []
    for i, u in enumerate(raw_units):
        rows.append({
            "id": f"{sid}:{i:04d}", "source_id": sid, "kind": u.kind, "ord": i, "text": u.text, "loc": u.loc,
            "label": u.label, "section": u.section, "image": u.image, "page_image": u.page_image,
            "topic_id": None, "subtopic_id": None, "concepts": [], "tokens": est_tokens(u.text),
        })
    store.insert_many("units", rows)
    progress(0.9, "Embedding units for retrieval")
    retrieval.invalidate()
    retrieval.get_index()
    from .. import topics  # local import to avoid a cycle

    if topics.has_topics():
        progress(0.95, "Tagging units to the course topic map")
        topics.tag_units([r["id"] for r in rows])
    meta["seconds"] = round(time.perf_counter() - t0, 1)
    store.update("sources", "id", sid, {"meta": {**meta, "short_title": short_title, "page_offset": page_offset, "units": len(rows)}})
    progress(1.0, f"Ingested {len(rows)} units from {title}")
    return sid


def _video_units(segs, short_title: str) -> list[RawUnit]:
    out = []
    for w in window_segments(segs):
        out.append(RawUnit(kind="transcript", text=w.text, loc={"start": round(w.start, 1), "end": round(w.end, 1)},
                           label=f"{short_title} @ {fmt_ts(w.start)}", section=""))
    return out
