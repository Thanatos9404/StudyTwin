from pathlib import Path

from studytwin.config import settings
from studytwin.ingest.video import Segment, fmt_ts, parse_caption_file, window_segments, youtube_id

DEMO = Path(__file__).resolve().parent.parent / "data" / "demo"


def test_slide_handout_is_split_into_individual_slides(kb):
    store = kb["store"]
    units = store.many("SELECT * FROM units WHERE source_id=? AND kind='slide' ORDER BY ord", (kb["ids"]["slides"],))
    assert [u["loc"]["slide"] for u in units] == list(range(1, 9))  # 2 handout pages x 4 slides
    assert units[2]["section"] == "Conditional probability"
    assert units[4]["label"] == "6.041 L2 slides · slide 5"
    for u in units:
        assert (settings.media_dir / u["image"]).exists()


def test_textbook_units_carry_printed_page_and_section(kb):
    store = kb["store"]
    units = store.many("SELECT * FROM units WHERE source_id=? ORDER BY ord", (kb["ids"]["textbook"],))
    assert any(u["kind"] == "figure" for u in units)
    rule = [u for u in units if "The Addition Rule" in u["text"]]
    assert rule and rule[0]["loc"]["page_label"] == 189
    assert rule[0]["section"].startswith("3.3")
    assert all(len(u["text"].split()) >= 8 for u in units if u["kind"] == "text")


def test_video_units_have_increasing_timestamps(kb):
    store = kb["store"]
    units = store.many("SELECT * FROM units WHERE source_id=? ORDER BY ord", (kb["ids"]["video"],))
    starts = [u["loc"]["start"] for u in units]
    assert starts == sorted(starts) and len(units) > 30
    assert all(u["loc"]["end"] - u["loc"]["start"] <= 95 for u in units)
    assert units[0]["label"].startswith("6.041 Lecture 2 @ 0:00")


def test_caption_parsing_and_windows():
    segs = parse_caption_file(DEMO / "MIT6_041F11_lec02_300k.srt")
    assert segs[0].start < 1.0 and segs[0].text.startswith("The following content")
    wins = window_segments([Segment(i * 5.0, i * 5.0 + 5, f"Sentence {i}.") for i in range(40)], target_s=60)
    assert all(w.end - w.start <= 95 for w in wins) and wins[0].start == 0.0
    assert fmt_ts(3725) == "1:02:05" and fmt_ts(75) == "1:15"
    assert youtube_id("https://www.youtube.com/watch?v=TluTv5V0RmE") == "TluTv5V0RmE"


def test_pptx_ingestion(tmp_path):
    from pptx import Presentation

    from studytwin.ingest.pipeline import ingest_pptx

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Bayes rule"
    s.placeholders[1].text = "P(A|B) = P(B|A)P(A)/P(B)"
    path = tmp_path / "deck.pptx"
    prs.save(path)
    units = ingest_pptx(path, tmp_path / "media", "deck", "Deck")
    assert units[0].kind == "slide" and units[0].loc == {"slide": 1} and "Bayes" in units[0].text
